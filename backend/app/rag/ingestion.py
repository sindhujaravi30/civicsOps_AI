
import csv
import hashlib
import io
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
DEFAULT_CHUNK_SIZE = 1200
DEFAULT_CHUNK_OVERLAP = 200

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".csv"}


@dataclass
class DocumentChunk:
    chunk_id: str
    text: str
    source: str
    file_type: str
    location: str
    chunk_index: int


def _decode_csv(content: bytes) -> str:
    """Decode common CSV encodings, including UTF-8 with a BOM."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError("Could not decode CSV. Please use UTF-8 encoding.")


def _extract_pdf(content: bytes) -> list[tuple[str, str]]:
    reader = PdfReader(io.BytesIO(content))

    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported.")

    sections = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append((f"page:{page_number}", text))

    return sections


def _extract_csv(content: bytes) -> list[tuple[str, str]]:
    text = _decode_csv(content)
    reader = csv.reader(io.StringIO(text))

    rows = list(reader)
    if not rows:
        return []

    headers = [value.strip() for value in rows[0]]
    sections = []

    for row_number, row in enumerate(rows[1:], start=2):
        values = []
        for index, value in enumerate(row):
            value = value.strip()
            if not value:
                continue

            header = (
                headers[index]
                if index < len(headers) and headers[index]
                else f"column_{index + 1}"
            )
            values.append(f"{header}: {value}")

        if values:
            sections.append((f"row:{row_number}", " | ".join(values)))

    # A CSV without a header still has useful content.
    if len(rows) == 1 and rows[0]:
        sections.append(("row:1", " | ".join(rows[0])))

    return sections


def _extract_excel(content: bytes) -> list[tuple[str, str]]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
    )

    sections = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                continue

            headers = [
                str(value).strip() if value is not None else ""
                for value in headers
            ]

            for row_number, row in enumerate(rows, start=2):
                values = []

                for index, value in enumerate(row):
                    if value is None or str(value).strip() == "":
                        continue

                    header = (
                        headers[index]
                        if index < len(headers) and headers[index]
                        else f"column_{index + 1}"
                    )
                    values.append(f"{header}: {str(value).strip()}")

                if values:
                    sections.append(
                        (
                            f"sheet:{worksheet.title},row:{row_number}",
                            " | ".join(values),
                        )
                    )
    finally:
        workbook.close()

    return sections


def _split_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """Split text into overlapping chunks without empty chunks."""
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be between zero and chunk_size - 1.")

    text = " ".join(text.split())
    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        # Prefer a word boundary when it is reasonably close.
        if end < len(text):
            boundary = text.rfind(" ", start, end)
            if boundary > start + chunk_size // 2:
                end = boundary

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        next_start = end - overlap
        start = next_start if next_start > start else end

    return chunks


def ingest_document(
    filename: str,
    content: bytes,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[DocumentChunk]:
    """Extract and chunk a supported document in memory."""
    safe_name = Path(filename).name
    extension = Path(safe_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload a PDF, XLSX, or CSV file."
        )

    if not content:
        raise ValueError("The uploaded file is empty.")

    if len(content) > MAX_FILE_BYTES:
        raise ValueError("File exceeds the 10 MB size limit.")

    if extension == ".pdf":
        sections = _extract_pdf(content)
    elif extension == ".xlsx":
        sections = _extract_excel(content)
    else:
        sections = _extract_csv(content)

    if not sections:
        raise ValueError(
            "No readable text or tabular data was found in this document."
        )

    chunks = []
    index = 0

    for location, section_text in sections:
        for chunk_text in _split_text(section_text, chunk_size, overlap):
            fingerprint = hashlib.sha256(
                f"{safe_name}|{location}|{index}|{chunk_text}".encode("utf-8")
            ).hexdigest()[:24]

            chunks.append(
                DocumentChunk(
                    chunk_id=fingerprint,
                    text=chunk_text,
                    source=safe_name,
                    file_type=extension.lstrip("."),
                    location=location,
                    chunk_index=index,
                )
            )
            index += 1

    return chunks
